"""Experiment 13 channel 1: ATO Individuals Table 6 by postcode, 2010-11 to 2022-23 -> NSW postcode x income-year
panel (PRESPEC 4.1). Uses the "all individuals by postcode" sheet of each file.
Run: uv run --no-project --with pandas --with openpyxl --with pyarrow python build_ato.py"""
import re
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

HERE = Path(__file__).resolve().parent
RAW = HERE / 'raw'
SHEET = {'2010-11': 'Individuals Tax Table 6a', '2011-12': 'Individual tax table 6A', '2012-13': 'Postcode only'}


def norm(s):
    s = str(s or '').replace('–', '-').replace('\n', ' ').lower()
    s = re.sub(r'(?<=[a-z])\d+', '', s)          # footnote digits glued to words
    return re.sub(r'\s+', ' ', s).strip()


def read_year(label):
    wb = openpyxl.load_workbook(RAW / f'ato_t6_{label}.xlsx', read_only=True)
    sheet = SHEET.get(label) or next(s for s in wb.sheetnames if s.strip().endswith('6B'))
    rows = list(wb[sheet].iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows) if any(norm(x) == 'postcode' for x in r if x))
    hdr = [norm(x) for x in rows[h]]
    first = h + 1
    if all(norm(x) in ('', 'no.', '$') for x in rows[h + 1]):      # separate units row (2010-11)
        hdr = [f'{a} {norm(b)}'.strip() for a, b in zip(hdr, rows[h + 1])]
        first = h + 2
    def col(pat, unit):
        for j, c in enumerate(hdr):
            if re.match(pat, c) and c.endswith(unit):
                return j
        raise KeyError((label, pat, unit))
    jp = hdr.index('postcode')
    js = next(j for j, c in enumerate(hdr) if c.startswith('state'))
    J = dict(individuals=next(j for j, c in enumerate(hdr) if re.match(r'^(number of )?individuals( no\.)?$', c)),
             taxinc=col(r'^taxable income or loss', '$'), sw_n=col(r'^salary or wages', 'no.'),
             sw_amt=col(r'^salary or wages', '$'),
             bus_pp=col(r'^total business income\W+primary production', '$'),
             bus_npp=col(r'^total business income\W+non\W*primary production', '$'))
    recs = []
    for r in rows[first:]:
        if r is None or len(r) <= max(J.values()) or r[js] != 'NSW':
            continue
        pc = pd.to_numeric(r[jp], errors='coerce')
        if pd.isna(pc):
            continue
        recs.append({'postcode': f'{int(pc):04d}', **{k: pd.to_numeric(r[j], errors='coerce') for k, j in J.items()}})
    d = pd.DataFrame(recs)
    d['year'] = int(label[:4])
    d['sheet'] = sheet
    d['header_used'] = '|'.join(f'{k}={hdr[j]}' for k, j in J.items())
    return d


def main():
    labels = [f'{y}-{str(y + 1)[2:]}' for y in range(2010, 2023)]
    d = pd.concat([read_year(l) for l in labels], ignore_index=True)
    d['bus_total'] = d[['bus_pp', 'bus_npp']].sum(axis=1, min_count=1)
    d.to_parquet(HERE / 'panel_ato.parquet', index=False)
    print(d.groupby('year').agg(n=('postcode', 'size'), sheet=('sheet', 'first')).to_string())
    print(d.drop_duplicates('year')[['year', 'header_used']].to_string())


if __name__ == '__main__':
    main()
