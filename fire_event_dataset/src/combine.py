"""One workbook for the mentor: key declared events first, then the full fire dataset.

Reads out/nsw_key_bushfire_events.xlsx and out/nsw_fire_events_2015_2025.xlsx (build those first) and writes
out/nsw_bushfires_2015_2025_combined.xlsx. URL columns are made clickable.
Run: uv run --no-sync --with openpyxl python -m src.combine
"""
import pandas as pd

from src.common import OUT

KEY = OUT / "nsw_key_bushfire_events.xlsx"
FULL = OUT / "nsw_fire_events_2015_2025.xlsx"
DEST = OUT / "nsw_bushfires_2015_2025_combined.xlsx"

# (new sheet name, source workbook, sheet, what it is)
SHEETS = [
    ("key_events", KEY, "events", "125 declared NSW bushfire disasters: summary, key facts, sources"),
    ("key_event_council", KEY, "event_council", "Each declared event × council: fire features, council context, Y columns"),
    ("key_facts", KEY, "facts", "Every reported fact with source, date, verbatim quote and link check"),
    ("key_dictionary", KEY, "dictionary", "Column meanings for the three key_ sheets"),
    ("key_sources", KEY, "sources", "Every document behind the key_ sheets"),
    ("all_fires", FULL, "fires", "All NSW bushfires 2015–2025, one row per fire × council, Bowen's template columns first"),
    ("lga_year", FULL, "lga_year", "Council × year economy and finance panel, 2014–2025"),
    ("all_fires_dictionary", FULL, "dictionary", "Column meanings, units, sources and coverage for all_fires"),
    ("lga_year_dictionary", FULL, "lga_year_dictionary", "Column sources for lga_year"),
    ("download_links", FULL, "download_links", "Exact download link for every input dataset"),
    ("data_sources", FULL, "sources", "Every input file: local path, SHA-256, access time, licence"),
]
WRAP_COLS = {"summary", "key_facts", "quoted_text", "notes", "what it holds", "meaning", "note", "main_fires",
             "fire_names", "councils", "towns_affected", "key_sources", "links_removed", "house_loss_all_sources",
             "official_declaration_name", "correction_note", "used_for", "dataset", "title"}
MIN_W, MAX_W, WRAP_W = 8, 45, 70  # column widths in characters


def format_sheet(ws, d):
    """Bold frozen header; width fitted to content (header and 95th-percentile cell length), long text wrapped."""
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    ws.freeze_panes = "B2" if len(d.columns) > 8 else "A2"
    bold, wrap = Font(bold=True), Alignment(wrap_text=True, vertical="top")
    for j, col in enumerate(d.columns, start=1):
        cell = ws.cell(row=1, column=j)
        cell.font = bold
        cell.alignment = Alignment(wrap_text=True, vertical="bottom")
        vals = d[col].dropna()
        # display formats only (stored values unchanged): dates as yyyy-mm-dd; large decimals with thousands separators
        fmt = None
        if pd.api.types.is_datetime64_any_dtype(d[col]):
            fmt = "yyyy-mm-dd"
        elif pd.api.types.is_float_dtype(d[col]) and len(vals):
            med = vals.abs().median()
            whole = bool((vals == vals.round()).all())
            fmt = "#,##0" if (med >= 100 or whole) else "#,##0.00" if med >= 1 else "0.0000"
            if any(k in str(col).lower() for k in ("lat", "lon")):
                fmt = "0.0000"  # coordinates: 4 decimals ≈ 10 m
        if fmt:
            for i in range(2, len(d) + 2):
                ws.cell(row=i, column=j).number_format = fmt
            ws.column_dimensions[get_column_letter(j)].width = max(12, min(len(str(col)), 30) + 2) if fmt == "yyyy-mm-dd" \
                else max(MIN_W, min(MAX_W, max(len(f"{vals.abs().max():,.0f}") + (6 if "." in fmt else 2),
                                                min(len(str(col)), 30) + 2)))
            continue
        vals = vals.map(lambda v: len(f"{v:,.2f}") if isinstance(v, float) else len(str(v)))
        body = int(vals.quantile(0.95)) if len(vals) else 0
        head = min(len(str(col)), 30)
        letter = get_column_letter(j)
        if col in WRAP_COLS and body > MAX_W:
            ws.column_dimensions[letter].width = WRAP_W
            for i in range(2, len(d) + 2):
                ws.cell(row=i, column=j).alignment = wrap
        else:
            ws.column_dimensions[letter].width = max(MIN_W, min(MAX_W, max(body, head) + 2))
    ws.row_dimensions[1].height = 30
    ws.auto_filter.ref = ws.dimensions


URL_COLS = {"source_url", "declaration_source_url", "download_url", "url", "source", "house_loss_url"}


def main():
    guide = pd.DataFrame(
        [(n, w) for n, _, _, w in SHEETS],
        columns=["sheet", "what it holds"])
    notes = pd.DataFrame({"sheet": ["", "Notes"], "what it holds": ["", ""]})
    rules = pd.DataFrame({"sheet": ["", "", "", "", ""], "what it holds": [
        "Blank means no data; 0 means the source was checked and the value is zero.",
        "Y, Y_class, DL, IL, FP, SL and split are left blank on purpose, to be designed in class.",
        "Every number traces to a source (see the dictionaries, key_sources, download_links); links were checked.",
        "Facts: a value is kept only if the number appears in its quote; hedged figures ('more than 700 ha') are blanked.",
        "_excess columns = change in the council minus the median change in councils without a fire of 100 ha or more.",
    ]})
    guide = pd.concat([guide, notes, rules], ignore_index=True)
    with pd.ExcelWriter(DEST, engine="openpyxl") as w:
        guide.to_excel(w, sheet_name="README", index=False)
        for name, book, sheet, _ in SHEETS:
            d = pd.read_excel(book, sheet_name=sheet)
            d.to_excel(w, sheet_name=name, index=False)
            ws = w.sheets[name]
            format_sheet(ws, d)
            for j, col in enumerate(d.columns, start=1):
                if col not in URL_COLS:
                    continue
                for i, u in enumerate(d[col], start=2):
                    if isinstance(u, str) and u.startswith("http") and "{" not in u and " " not in u:
                        ws.cell(row=i, column=j).hyperlink = u
                        ws.cell(row=i, column=j).style = "Hyperlink"
        g = w.sheets["README"]
        format_sheet(g, guide)
        g.column_dimensions["A"].width = 24
        g.column_dimensions["B"].width = 110
        g.auto_filter.ref = None
    return DEST


if __name__ == "__main__":
    print(main())
