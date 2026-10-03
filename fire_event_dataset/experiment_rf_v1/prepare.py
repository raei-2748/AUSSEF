"""Freeze the user-selected workbook and audit its master sheet without fitting models."""
import argparse
import csv
import hashlib
import json
import shutil
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
SOURCE = Path('/Users/ray/Research/AUSSEF - Local/data/master_workbook/nsw_bushfires_2015_2025_XY.xlsx')


def digest(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def sheet_frame(ws, header_row=1):
    it = ws.iter_rows(min_row=header_row, values_only=True)
    headers = list(next(it))
    assert all(headers) and len(set(headers)) == len(headers), ws.title
    return pd.DataFrame(it, columns=headers)


def prepare(source=SOURCE):
    snap = ROOT / 'snapshot'
    snap.mkdir(parents=True, exist_ok=True)
    dest = snap / source.name
    source_hash = digest(source)
    if dest.exists():
        assert digest(dest) == source_hash, 'Source changed; use a new experiment rather than replacing V1.'
    else:
        shutil.copyfile(source, dest)
    w = load_workbook(dest, read_only=True, data_only=True)
    m = sheet_frame(w['master'], 3)
    cb = sheet_frame(w['codebook'])
    rm = sheet_frame(w['README'])
    master_cb = cb[cb.sheet == 'master'].set_index('column')
    assert len(m) == 218 and not m.duplicated(['agrn', 'region_id']).any()
    assert set(m.columns) == set(master_cb.index)
    assert pd.to_numeric(m.Y, errors='coerce').notna().all()
    assert m.Y.between(0, 1).all()
    assert np.allclose(m[['DL', 'IL', 'FP', 'SL']].mean(axis=1), m.Y)
    rank = m.Y.rank(pct=True)
    classes = np.select([rank >= .95, rank >= .8, rank >= .5], [4, 3, 2], 1)
    assert (classes == m.Y_class_from_Y).all()
    cuts = [float(m.loc[m.Y_class_from_Y == k, 'Y'].min()) for k in [2, 3, 4]]
    assert (np.digitize(m.Y, cuts, right=False) + 1 == m.Y_class_from_Y).all()
    m['row_id'] = m.agrn.astype(str) + '|' + m.region_id.astype(str)
    m['event_start'] = pd.to_datetime(m.first_fire_start)
    m['fire_fy_start'] = np.where(m.event_start.dt.month >= 7, m.event_start.dt.year, m.event_start.dt.year - 1)
    # Keep declarations, shared mapped fires, and reused council-FY outcomes together.
    parent = list(range(len(m)))

    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    def union(a, b):
        a, b = find(a), find(b)
        if a != b:
            parent[max(a, b)] = min(a, b)

    seen = {}
    reasons = []
    for i, r in m.iterrows():
        keys = [('declaration', str(r.agrn)), ('council_fy', str(r.region_id), int(r.fire_fy_start))]
        keys += [('fire', x.strip()) for x in str(r.info_fire_event_ids).split(';') if x.strip() and x.strip() != 'None']
        for key in keys:
            if key in seen:
                union(i, seen[key])
                reasons.append({'row_a': m.row_id[i], 'row_b': m.row_id[seen[key]], 'reason': '|'.join(map(str, key))})
            else:
                seen[key] = i
    # Explicit overlapping declarations also bind even if their linked mapped fire lists differ.
    by_agrn = {str(a): list(g.index) for a, g in m.groupby('agrn')}
    for i, r in m.iterrows():
        for a in str(r.info_also_under_declarations).split(';'):
            if a.strip() in by_agrn:
                j = by_agrn[a.strip()][0]
                union(i, j)
                reasons.append({'row_a': m.row_id[i], 'row_b': m.row_id[j], 'reason': 'overlap_declaration|' + a.strip()})
    roots = sorted({find(i) for i in range(len(m))})
    ids = {r: f'E{j + 1:03d}' for j, r in enumerate(roots)}
    m['event_group'] = [ids[find(i)] for i in range(len(m))]
    g = m.groupby('event_group').agg(rows=('row_id', 'size'), declarations=('agrn', 'nunique'),
                                    first_start=('event_start', 'min'), last_start=('event_start', 'max'),
                                    fy_min=('fire_fy_start', 'min'), fy_max=('fire_fy_start', 'max'))
    for name, frame in [('master', m), ('codebook', cb), ('readme', rm), ('group_edges', pd.DataFrame(reasons)), ('event_groups', g.reset_index())]:
        frame.to_csv(snap / f'{name}.csv', index=False)
    candidates = master_cb[master_cb.role_if_Y_sum == 'X'].index.tolist()
    audit_rows = []
    for c in candidates:
        x = m[c].replace('N/A', np.nan)
        number = pd.to_numeric(x, errors='coerce')
        numerical = (number.notna().sum() == x.notna().sum())
        audit_rows.append({'column': c, 'code': master_cb.loc[c, 'code'], 'meaning': master_cb.loc[c, 'meaning'],
                           'coverage': float(x.notna().mean()), 'unique': int(x.nunique()),
                           'type': 'numeric' if numerical else 'categorical'})
    pd.DataFrame(audit_rows).to_csv(snap / 'feature_audit.csv', index=False)
    audit = {'source': str(source), 'source_sha256': source_hash, 'snapshot_sha256': digest(dest),
             'frozen_at_utc': datetime.now(timezone.utc).isoformat(), 'rows': len(m),
             'declarations': int(m.agrn.nunique()), 'councils': int(m.region_id.nunique()),
             'event_groups': len(g), 'largest_groups': g.sort_values('rows', ascending=False).head(8).reset_index().astype(str).to_dict('records'),
             'years': m.event_start.dt.year.value_counts().sort_index().to_dict(),
             'fy_counts': m.fire_fy_start.value_counts().sort_index().to_dict(),
             'pillars_n': m.Y_pillars_n.value_counts().sort_index().to_dict(),
             'class_from_y_counts': m.Y_class_from_Y.value_counts().sort_index().to_dict(),
             'final_class_floor_changes': int((m.Y_class != m.Y_class_from_Y).sum()),
             'reference_y_class_cutoffs': cuts, 'X_columns': len(candidates),
             'target_status': 'Frozen workbook target, ranked across the full reference cohort; not independent measurement validation.'}
    (ROOT / 'AUDIT.json').write_text(json.dumps(audit, indent=2, default=int) + '\n')
    assert digest(source) == source_hash
    w.close()
    print(json.dumps(audit, indent=2, default=int))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--source', type=Path, default=SOURCE)
    prepare(parser.parse_args().source)
