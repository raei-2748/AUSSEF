"""Independently verify saved predictions and produce the experiment handoff."""
import hashlib
import json
import subprocess
import sys
from collections import defaultdict
from pathlib import Path

import joblib
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.metrics import confusion_matrix, mean_absolute_error, r2_score

import run_experiment as runner

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'results'


def digest(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def main():
    audit = json.loads((ROOT/'AUDIT.json').read_text())
    receipt = json.loads((OUT/'RUN_RECEIPT.json').read_text())
    summary = json.loads((OUT/'SUMMARY.json').read_text())
    m = pd.read_csv(ROOT/'snapshot/master.csv',dtype={'agrn':str,'region_id':str})
    pred = pd.read_csv(OUT/'HELD_OUT_PREDICTIONS.csv',dtype={'agrn':str,'region_id':str})
    splits = pd.read_csv(OUT/'SPLIT_ASSIGNMENTS.csv')
    features = pd.read_csv(OUT/'FEATURE_DECISIONS.csv')
    cb = pd.read_csv(ROOT/'snapshot/codebook.csv').query("sheet == 'master'").set_index('column')
    checks = {}
    checks['drive_and_snapshot_hash_unchanged'] = digest(Path(audit['source'])) == digest(ROOT/'snapshot/nsw_bushfires_2015_2025_XY.xlsx') == audit['source_sha256']
    checks['protocol_unchanged_since_before_fit'] = digest(ROOT/'PROTOCOL.md') == receipt['protocol_sha256']
    checks['runner_unchanged_since_before_fit'] = digest(ROOT/'run_experiment.py') == receipt['runner_sha256']
    checks['splits_unchanged_since_before_fit'] = digest(OUT/'SPLIT_ASSIGNMENTS.csv') == receipt['split_sha256']
    # Read the actual supplied Excel Y values again; no dependence on preparation code.
    w=load_workbook(Path(audit['source']),read_only=True,data_only=True)
    it=w['master'].iter_rows(min_row=3,values_only=True);h=next(it);raw=pd.DataFrame(it,columns=h);w.close()
    keys = ['agrn','region_id']
    raw['row_id']=raw.agrn.astype(str)+'|'+raw.region_id.astype(str)
    raw=raw.set_index('row_id');mm=m.set_index('row_id')
    checks['all_218_rows_preserved_with_unique_keys'] = len(m)==218 and m.row_id.nunique()==218 and set(m.row_id)==set(raw.index)
    checks['Y_exactly_preserved_from_Drive_master'] = bool(np.allclose(raw.Y.reindex(mm.index),mm.Y,rtol=0,atol=1e-15))
    checks['Y_is_available_pillar_mean'] = bool(np.allclose(mm[['DL','IL','FP','SL']].mean(axis=1),mm.Y,atol=1e-15))
    cuts=np.array(summary['reference_class_cutoffs'])
    checks['frozen_cutoffs_reproduce_workbook_class_from_Y'] = bool((np.digitize(mm.Y,cuts)+1==mm.Y_class_from_Y).all())
    memberships=defaultdict(set)
    for _,r in m.iterrows():
        memberships[('declaration',r.agrn)].add(r.event_group)
        memberships[('council_fy',r.region_id,r.fire_fy_start)].add(r.event_group)
        for eid in str(r.info_fire_event_ids).split(';'):
            if eid.strip() and eid.strip()!='nan':memberships[('fire',eid.strip())].add(r.event_group)
    checks['shared_declarations_fires_and_council_years_grouped'] = all(len(v)==1 for v in memberships.values())
    all_disjoint=True
    no_ids=True
    for name,s in splits.groupby('split'):
        tr=s[s.role=='train'];te=s[s.role=='test']
        all_disjoint &= set(tr.event_group).isdisjoint(te.event_group)
        no_ids &= set(tr.row_id).isdisjoint(te.row_id)
    checks['every_outer_split_has_disjoint_event_groups']=bool(all_disjoint)
    checks['every_outer_split_has_disjoint_rows']=bool(no_ids)
    pri=splits[splits.split=='primary_recent_seasons']
    train=mm.loc[pri[pri.role=='train'].row_id]
    test=mm.loc[pri[pri.role=='test'].row_id]
    checks['primary_is_139_old_vs_79_recent_rows']=len(train)==139 and len(test)==79
    checks['primary_temporal_embargo_satisfied']=bool(train.fire_fy_start.max()<=2019 and test.fire_fy_start.min()>=2022)
    last_outcome=pd.Timestamp(int(train.fire_fy_start.max())+2,6,30)
    checks['latest_training_t_plus1_FY_end_precedes_first_test_fire']=bool(last_outcome<pd.to_datetime(test.first_fire_start).min())
    cv=pred[pred.split.str.startswith('group_cv_')]
    checks['every_row_has_one_grouped_out_of_fold_prediction']=len(cv)==218 and cv.row_id.nunique()==218
    checks['predictions_finite_and_in_Y_support']=bool(np.isfinite(pred[['pred_mean','pred_median','pred_ridge','pred_rf']]).all().all() and pred[['pred_mean','pred_median','pred_ridge','pred_rf']].ge(0).all().all() and pred[['pred_mean','pred_median','pred_ridge','pred_rf']].le(1).all().all())
    checks['predicted_classes_come_only_from_continuous_prediction']=bool((np.digitize(pred.pred_rf,cuts)+1==pred.pred_class_rf).all())
    retained=features[features.decision=='retained'].column.unique()
    checks['only_codebook_X_columns_used_as_predictors']=all(cb.loc[c,'role_if_Y_sum']=='X' for c in retained)
    checks['no_Y_components_classes_or_pillar_missingness_predictors']=all(not c.startswith(('Y','DL','IL','FP','SL','info_')) for c in retained)
    checks['full_future_reference_growth_columns_excluded']=not any(c in retained for c in ['X_socio_grp_sa4_nominal_growth_pct','X_socio_grp_sa4_real_growth_pct'])
    # Recompute the simplest constant from the training group means independently.
    constants_correct=True
    for name,s in splits.groupby('split'):
        tr=mm.loc[s[s.role=='train'].row_id]
        want=tr.groupby('event_group').Y.mean().mean()
        got=pred.loc[pred.split==name,'pred_mean']
        constants_correct &= bool(np.allclose(got,want,atol=1e-15))
    checks['constant_mean_baselines_use_training_groups_only']=bool(constants_correct)
    # Export a standard sklearn object bundle that loads in a fresh Python process.
    sys.modules['__main__'].PreparedModel=runner.PreparedModel
    custom=joblib.load(OUT/'PRIMARY_RF_MODEL.joblib')
    model=custom['model']
    safe={'pipeline':model.pipeline,'numeric_columns':model.numeric,'categorical_columns':model.categorical,
          'columns':model.columns,'class_cutoffs':cuts.tolist(),'source_sha256':audit['source_sha256'],
          'training_row_ids':custom['training_row_ids'],'status':custom['status'],
          'prediction_timing':'retrospective after final fire characteristics; not certified real-time',
          'availability_rules_file':'AVAILABILITY_RULES.csv'}
    joblib.dump(safe,OUT/'RF_MODEL.joblib')
    x,_=runner.mask_inputs(m,cb.reset_index())
    primary=pred[pred.split=='primary_recent_seasons']
    ii=[int(m.index[m.row_id==rid][0]) for rid in primary.row_id]
    checks['saved_model_reproduces_held_out_predictions']=bool(np.allclose(model.predict(x.iloc[ii]),primary.pred_rf,atol=1e-14))
    checks['bootstrap_has_all_three_baseline_comparisons']=len(pd.read_csv(OUT/'PAIRED_GROUP_BOOTSTRAP.csv'))==6
    saved_metrics=pd.read_csv(OUT/'MODEL_RESULTS.csv')
    metric_ok=True
    for name,d in [('primary_recent_seasons',primary),('group_cv_pooled',cv)]:
        for model_name in ['mean','median','ridge','rf']:
            row=saved_metrics[(saved_metrics.split==name)&(saved_metrics.model==model_name)].iloc[0]
            errors=np.abs(d.Y.to_numpy()-d[f'pred_{model_name}'].to_numpy())
            by_group=[]
            for group in sorted(d.event_group.unique()):
                by_group.append(float(errors[d.event_group.to_numpy()==group].mean()))
            metric_ok &= bool(np.isclose(np.mean(by_group),row.group_mae,atol=1e-14) and np.isclose(errors.mean(),row.row_mae,atol=1e-14))
    checks['reported_group_and_row_errors_match_saved_predictions']=bool(metric_ok)
    fresh=subprocess.run([sys.executable,'-c',
                          "import joblib,sys; b=joblib.load(sys.argv[1]); assert 'pipeline' in b and len(b['columns'])>0",str(OUT/'RF_MODEL.joblib')],
                         capture_output=True,text=True)
    checks['portable_model_loads_in_fresh_python_process']=fresh.returncode==0
    result={'status':'PASS' if all(checks.values()) else 'FAIL','checks':checks,'passed':sum(checks.values()),'total':len(checks)}
    (OUT/'VALIDATION_CHECKS.json').write_text(json.dumps(result,indent=2)+'\n')
    assert all(checks.values()),result
    report(m,pred,summary,audit,result)
    files={str(p.relative_to(ROOT)):digest(p) for p in ROOT.rglob('*') if p.is_file() and '.venv' not in p.parts and '.mplconfig' not in p.parts and '__pycache__' not in p.parts and p.name!='OUTPUT_HASHES.json'}
    (OUT/'OUTPUT_HASHES.json').write_text(json.dumps(files,indent=2)+'\n')
    print(json.dumps(result,indent=2))


def table(metrics):
    lines=['| Model | Group MAE | Row MAE | RMSE | R² | Class accuracy | Macro-F1 |',
           '|---|---:|---:|---:|---:|---:|---:|']
    for r in metrics:
        lines.append(f"| {r['model']} | {r['group_mae']:.4f} | {r['row_mae']:.4f} | {r['rmse']:.4f} | {r['r2']:.3f} | {r['class_accuracy']:.1%} | {r['class_macro_f1']:.3f} |")
    return '\n'.join(lines)


def report(m,pred,s,audit,validation):
    verdict='The frozen experiment found transferable predictive gain on this workbook benchmark.' if s['verdict'].startswith('PREDICTIVE_GAIN_ESTABLISHED') else 'The frozen experiment did not establish transferable predictive gain over the simple baselines.'
    selections=pd.read_csv(OUT/'SELECTED_MODELS.csv')
    selected=selections[(selections.split=='primary_recent_seasons')&(selections.model=='rf')].iloc[0]
    primary=pred[pred.split=='primary_recent_seasons']
    cv=pred[pred.split.str.startswith('group_cv_')]
    p_rf=next(x for x in s['primary_metrics'] if x['model']=='rf')
    p_best=min((x for x in s['primary_metrics'] if x['model']!='rf'),key=lambda x:x['group_mae'])
    improved=(p_best['group_mae']-p_rf['group_mae'])/p_best['group_mae']*100
    im=pd.read_csv(OUT/'PERMUTATION_IMPORTANCE.csv').head(8)
    importance='\n'.join(f"- `{r.column}`: held-out group MAE increase {r.group_mae_increase:.4f}." for _,r in im.iterrows())
    bootstrap='\n'.join(f"- Versus {r['baseline']}: {r['improvement_pct']:+.1f}% improvement; group-MAE difference {r['improvement_group_mae']:+.4f}, conditional 95% interval [{r['ci_low']:+.4f}, {r['ci_high']:+.4f}]." for r in s['primary_bootstrap'])
    recall=[]
    for name,d in [('Recent-season holdout',primary),('Grouped cross-validation',cv)]:
        mat=confusion_matrix(d.Y_class_from_Y,d.pred_class_rf,labels=[1,2,3,4])
        for k in range(4):
            n=int(mat[k].sum());recall.append(f'| {name} | {k+1} | {n} | {mat[k,k]/n:.1%} |' if n else f'| {name} | {k+1} | 0 | N/A |')
    known=set(m.loc[m.fire_fy_start<=2019,'region_id'])
    unseen=int((~primary.region_id.isin(known)).sum())
    group_details=pd.read_csv(ROOT/'snapshot/event_groups.csv').sort_values('rows',ascending=False).head(3)
    cohort=[]
    for label,part in [('Older training rows',m[m.fire_fy_start<=2019]),('Recent held-out rows',m[m.fire_fy_start>=2022])]:
        cohort.append(f"| {label} | {len(part)} | {part.Y.median():.3f} | {(part.Y_pillars_n==4).sum()} | {(part.Y_class_from_Y>=3).sum()} |")
    text=f'''# Bowen random-forest experiment V1

Completed 28 September 2026. Input: the user-selected Drive workbook, `master` sheet.

{verdict} On the recent-season holdout, RF group MAE was **{p_rf['group_mae']:.4f}**, versus **{p_best['group_mae']:.4f}** for the strongest baseline ({p_best['model']}); the RF improvement was **{improved:+.1f}%**. This is a result for the supplied reference-cohort Y, not a claim about all fires or independently validated economic loss.

## What was tested

One row is a declared event × council. The 218 rows cover 96 declarations and 69 councils, but shared fires, overlapping declarations and reused council-FY outcomes join them into 49 connected groups. The largest is 67 rows and spans FY2018–FY2019, including Black Summer. These are dependency groups, not 49 certified independent fire episodes. There are no event-start rows in calendar years 2020–2022 in this master view; the results must not be advertised as uniform coverage of every fire in 2015–2025.

Y is preserved byte-for-byte numerically from the workbook. Its core indicators are homes lost per 1,000 dwellings (DL); excess total personal-income and business-count falls (IL); excess cash-cover drawdown, services crowd-out and renewals-ratio rise (FP); excess income-support-recipient rise (SL). Each is percentile-ranked, ranks are averaged into pillars, and available pillars are averaged into Y. 84 rows have four pillars, 116 have three and 18 have two. Missing outcomes were not imputed. This changes the effective weights and is part of the existing definition.

The regression predicts Y. Predicted classes come from the fixed reference boundaries **0.486631, 0.635393, 0.754397**. They reproduce `Y_class_from_Y`. The separate `Y_class` includes death/home-loss floors and differs in 23 rows; those outcome-derived overrides are not predicted or applied.

## Validation design

The primary temporal holdout trains on 139 rows from complete groups ending no later than FY2019 and tests 79 rows from groups beginning in FY2022 or later. The latest training t+1 outcome FY ends 30 June 2021. No held-out declaration, linked mapped fire or council-FY outcome group appears in training. There are {unseen} test rows from councils absent from training; the remaining test councils recur. This tests new events, not exclusively new councils.

The secondary five-fold grouped cross-validation predicts every row once out of fold. It is a cross-event benchmark with no time ordering. In both designs, each outer training set uses four-fold grouped inner tuning. All numerical imputation, categorical encoding, constant/duplicate removal and coverage screening are learned from training cells only. Model fitting and primary MAE give each connected group equal total weight. Row metrics are also reported.

Candidates were the 223 codebook X columns. Two long-window GRP-growth predictors were excluded. Later source reference years were masked only in the experiment copy. Within training, a column needed >=50% coverage and nonconstant values; exact duplicates were dropped. The primary forest retained **{int(selected.retained_features)}** original features. It uses 500 trees and the frozen selected settings `{selected.parameters}`. Predictor IDs, Y ingredients, impact-side columns, class labels and pillar availability were excluded. The forest is a retrospective post-fire estimator; final area and mapped severity were available as inputs. It is not an early-warning system.

## Recent seasons held out

{table(s['primary_metrics'])}

Positive differences below favour RF. Intervals resample whole held-out groups and condition on the already-fitted models; they do not represent all model-selection uncertainty.

{bootstrap}

## All rows: grouped cross-validation

{table(s['cv_metrics'])}

![Model comparison](results/model_comparison.png)

![Held-out predictions](results/predicted_vs_actual.png)

## Classes derived from predicted Y

These are the four reference-relative classes. Extreme-class recall remains visible even when poor.

| Test | Actual class | Rows | RF recall |
|---|---:|---:|---:|
{chr(10).join(recall)}

Full confusion counts are in `results/CLASS_CONFUSION.csv`. Regression is primary; class conversion is secondary. Thresholds were fixed from the supplied workbook before fitting and were not optimised on prediction accuracy.

## Exploratory diagnosis after the frozen test

The primary RF predictions occupy only **{primary.pred_rf.min():.3f}–{primary.pred_rf.max():.3f}**, although actual held-out Y occupies **{primary.Y.min():.3f}–{primary.Y.max():.3f}**. Its highest prediction is below the severe-class boundary, so it misses every severe and extreme row. Ridge's range is wider, but its held-out R² is also negative. The data have not established a dependable recent-season predictor.

| Cohort | Rows | Median Y | Rows with four pillars | Severe/extreme from Y |
|---|---:|---:|---:|---:|
{chr(10).join(cohort)}

This is an exploratory explanation check, not an additional model or an attribution of cause. The recent cohort has a higher target distribution and much less four-pillar coverage. Missing-pillar reweighting may contribute to differences between cohorts, but this table does not establish that it caused the prediction failure. `COHORT_DIAGNOSTICS.csv` gives the FY detail. Review the comparability and interpretation of Y with Bowen before choosing a new model or changing the target.

## Interpretation

The prespecified success rule requires at least 5% primary group-MAE gain and a positive lower bootstrap improvement bound versus **all three** baselines, plus a lower grouped-CV MAE versus all three. Verdict: **{s['verdict']}**. No model grid, target or subset was changed after seeing predictive results. A no-gain result is about this benchmark and design; it does not imply that fires have no socioeconomic effect.

Highest descriptive permutation scores on the primary held-out set:

{importance}

These scores describe predictive reliance, not causal mechanisms or validated policy priorities. Correlated variables can substitute for one another. Row-wise permutations can produce implausible combinations and retain repeated group structure; use these only as exploratory interpretation. They were not used for model selection.

## Limits that remain

- Labels were constructed using percentile ranks over the full workbook reference cohort. They were preserved by explicit user choice. This is a frozen-index benchmark, not a fully forward-built target or independent validation of Y.
- Pillar coverage varies, with homes lost known for only 90 rows. Sensitivity by two/three/four pillars is saved in `ROBUSTNESS_RESULTS.csv`; it does not redefine Y.
- Annual LGA outcomes can dilute fire effects and repeat across events. Shared comparison benchmarks, temporal overlap of adjacent years and statewide shocks can leave dependence beyond the grouping used here.
- The oldest and newest cohorts have different source coverage. Source reference years are checked, but exact historical publication dates and later revisions are not certified. Future-looking benchmark performance is therefore conditional on retrospective measurements, not a proven operational forecast.
- The sample is a selected declared-event view with mapped-fire links, not the entire official inventory. There are 125 declarations in the workbook inventory, 96 represented in the master, and 6,146 fires in the separate inventory.
- Group bootstrap intervals are conditional and based on a limited number of groups. The workbook index is relative severity, not dollars and not a causal effect.

## Reproduction and saved outputs

The frozen workbook hash is `{audit['source_sha256']}`. Before-fit protocol, source, runner and split hashes are in `PRE_FIT_RECEIPT.json`; runtime versions and completion status are in `RUN_RECEIPT.json`. Validation: **{validation['status']} — {validation['passed']}/{validation['total']} checks**. The original Drive file was unchanged.

`results/HELD_OUT_PREDICTIONS.csv` contains actual Y, each model's prediction, RF-derived class and group for every test row. `MODEL_RESULTS.csv`, `PAIRED_GROUP_BOOTSTRAP.csv`, `INNER_TUNING.csv`, `SELECTED_MODELS.csv`, `FEATURE_DECISIONS.csv`, `AVAILABILITY_RULES.csv`, `ROBUSTNESS_RESULTS.csv` and `CLASS_CONFUSION.csv` expose the full experiment. `RF_MODEL.joblib` is a standard sklearn primary-training model bundle. It is labelled experimental and has not been refitted on the test rows.

Recreate an isolated Python environment with `requirements.txt`, run `prepare.py` on a fresh experiment copy, then `run_experiment.py` and `validate_and_report.py`. The runner refuses to overwrite a completed V1 receipt. No canonical database is a model input.

Methods: [random forest regression](https://scikit-learn.org/stable/modules/generated/sklearn.ensemble.RandomForestRegressor.html), [grouped validation](https://scikit-learn.org/stable/modules/cross_validation.html), [training-only preprocessing](https://scikit-learn.org/stable/common_pitfalls.html).
'''
    (ROOT/'REPORT.md').write_text(text)


if __name__=='__main__':
    main()
