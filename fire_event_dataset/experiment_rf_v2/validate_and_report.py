"""Independently verify saved predictions and produce the experiment handoff."""
import hashlib
import json
import subprocess
import sys
from collections import defaultdict
from pathlib import Path

import joblib
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.metrics import confusion_matrix, mean_absolute_error, r2_score

import run_experiment as runner

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'results'


def digest(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def main():
    audit = json.loads((ROOT/'AUDIT.json').read_text())
    receipt = json.loads((OUT/'RUN_RECEIPT.json').read_text())
    summary = json.loads((OUT/'SUMMARY.json').read_text())
    m = pd.read_csv(ROOT/'snapshot/master.csv',dtype={'agrn':str,'region_id':str})
    pred = pd.read_csv(OUT/'HELD_OUT_PREDICTIONS.csv',dtype={'agrn':str,'region_id':str})
    splits = pd.read_csv(OUT/'SPLIT_ASSIGNMENTS.csv')
    features = pd.read_csv(OUT/'FEATURE_DECISIONS.csv')
    cb = pd.read_csv(ROOT/'snapshot/codebook.csv').query("sheet == 'master'").set_index('column')
    checks = {}
    checks['drive_and_snapshot_hash_unchanged'] = digest(Path(audit['source'])) == digest(ROOT/'snapshot/nsw_bushfires_2015_2025_XY.xlsx') == audit['source_sha256']
    checks['protocol_unchanged_since_before_fit'] = digest(ROOT/'PROTOCOL.md') == receipt['protocol_sha256']
    checks['runner_unchanged_since_before_fit'] = digest(ROOT/'run_experiment.py') == receipt['runner_sha256']
    checks['splits_unchanged_since_before_fit'] = digest(OUT/'SPLIT_ASSIGNMENTS.csv') == receipt['split_sha256']
    # Read the actual supplied Excel Y values again; no dependence on preparation code.
    w=load_workbook(Path(audit['source']),read_only=True,data_only=True)
    it=w['master'].iter_rows(min_row=3,values_only=True);h=next(it);raw=pd.DataFrame(it,columns=h);w.close()
    keys = ['agrn','region_id']
    raw['row_id']=raw.agrn.astype(str)+'|'+raw.region_id.astype(str)
    raw=raw.set_index('row_id');mm=m.set_index('row_id')
    checks['all_218_rows_preserved_with_unique_keys'] = len(m)==218 and m.row_id.nunique()==218 and set(m.row_id)==set(raw.index)
    checks['Y_exactly_preserved_from_Drive_master'] = bool(np.allclose(raw.Y.reindex(mm.index),mm.Y,rtol=0,atol=1e-15))
    checks['Y_is_available_pillar_mean'] = bool(np.allclose(mm[['DL','IL','FP','SL']].mean(axis=1),mm.Y,atol=1e-15))
    cuts=np.array(summary['reference_class_cutoffs'])
    checks['frozen_cutoffs_reproduce_workbook_class_from_Y'] = bool((np.digitize(mm.Y,cuts)+1==mm.Y_class_from_Y).all())
    memberships=defaultdict(set)
    for _,r in m.iterrows():
        memberships[('declaration',r.agrn)].add(r.event_group)
        memberships[('council_fy',r.region_id,r.fire_fy_start)].add(r.event_group)
        for eid in str(r.info_fire_event_ids).split(';'):
            if eid.strip() and eid.strip()!='nan':memberships[('fire',eid.strip())].add(r.event_group)
    checks['shared_declarations_fires_and_council_years_grouped'] = all(len(v)==1 for v in memberships.values())
    all_disjoint=True
    no_ids=True
    for name,s in splits.groupby('split'):
        tr=s[s.role=='train'];te=s[s.role=='test']
        all_disjoint &= set(tr.event_group).isdisjoint(te.event_group)
        no_ids &= set(tr.row_id).isdisjoint(te.row_id)
    checks['every_outer_split_has_disjoint_event_groups']=bool(all_disjoint)
    checks['every_outer_split_has_disjoint_rows']=bool(no_ids)
    pri=splits[splits.split=='primary_recent_seasons']
    train=mm.loc[pri[pri.role=='train'].row_id]
    test=mm.loc[pri[pri.role=='test'].row_id]
    checks['primary_is_139_old_vs_79_recent_rows']=len(train)==139 and len(test)==79
    checks['primary_temporal_embargo_satisfied']=bool(train.fire_fy_start.max()<=2019 and test.fire_fy_start.min()>=2022)
    last_outcome=pd.Timestamp(int(train.fire_fy_start.max())+2,6,30)
    checks['latest_training_t_plus1_FY_end_precedes_first_test_fire']=bool(last_outcome<pd.to_datetime(test.first_fire_start).min())
    cv=pred[pred.split.str.startswith('group_cv_')]
    checks['every_row_has_one_grouped_out_of_fold_prediction']=len(cv)==218 and cv.row_id.nunique()==218
    checks['predictions_finite_and_in_Y_support']=bool(np.isfinite(pred[['pred_mean','pred_median','pred_ridge','pred_rf']]).all().all() and pred[['pred_mean','pred_median','pred_ridge','pred_rf']].ge(0).all().all() and pred[['pred_mean','pred_median','pred_ridge','pred_rf']].le(1).all().all())
    checks['predicted_classes_come_only_from_continuous_prediction']=bool((np.digitize(pred.pred_rf,cuts)+1==pred.pred_class_rf).all())
    retained=features[features.decision=='retained'].column.unique()
    checks['only_codebook_X_columns_used_as_predictors']=all(cb.loc[c,'role_if_Y_sum']=='X' for c in retained)
    checks['no_Y_components_classes_or_pillar_missingness_predictors']=all(not c.startswith(('Y','DL','IL','FP','SL','info_')) for c in retained)
    checks['full_future_reference_growth_columns_excluded']=not any(c in retained for c in ['X_socio_grp_sa4_nominal_growth_pct','X_socio_grp_sa4_real_growth_pct'])
    # Recompute the simplest constant from the training group means independently.
    constants_correct=True
    for name,s in splits.groupby('split'):
        tr=mm.loc[s[s.role=='train'].row_id]
        want=tr.groupby('event_group').Y.mean().mean()
        got=pred.loc[pred.split==name,'pred_mean']
        constants_correct &= bool(np.allclose(got,want,atol=1e-15))
    checks['constant_mean_baselines_use_training_groups_only']=bool(constants_correct)
    # Export a standard sklearn object bundle that loads in a fresh Python process.
    sys.modules['__main__'].PreparedModel=runner.PreparedModel
    custom=joblib.load(OUT/'PRIMARY_RF_MODEL.joblib')
    model=custom['model']
    safe={'pipeline':model.pipeline,'numeric_columns':model.numeric,'categorical_columns':model.categorical,
          'columns':model.columns,'class_cutoffs':cuts.tolist(),'source_sha256':audit['source_sha256'],
          'training_row_ids':custom['training_row_ids'],'status':custom['status'],
          'prediction_timing':'retrospective after final fire characteristics; not certified real-time',
          'availability_rules_file':'AVAILABILITY_RULES.csv'}
    joblib.dump(safe,OUT/'RF_MODEL.joblib')
    x,_=runner.mask_inputs(m,cb.reset_index())
    primary=pred[pred.split=='primary_recent_seasons']
    ii=[int(m.index[m.row_id==rid][0]) for rid in primary.row_id]
    checks['saved_model_reproduces_held_out_predictions']=bool(np.allclose(model.predict(x.iloc[ii]),primary.pred_rf,atol=1e-14))
    checks['bootstrap_has_all_three_baseline_comparisons']=len(pd.read_csv(OUT/'PAIRED_GROUP_BOOTSTRAP.csv'))==6
    saved_metrics=pd.read_csv(OUT/'MODEL_RESULTS.csv')
    metric_ok=True
    for name,d in [('primary_recent_seasons',primary),('group_cv_pooled',cv)]:
        for model_name in ['mean','median','ridge','rf']:
            row=saved_metrics[(saved_metrics.split==name)&(saved_metrics.model==model_name)].iloc[0]
            errors=np.abs(d.Y.to_numpy()-d[f'pred_{model_name}'].to_numpy())
            by_group=[]
            for group in sorted(d.event_group.unique()):
                by_group.append(float(errors[d.event_group.to_numpy()==group].mean()))
            metric_ok &= bool(np.isclose(np.mean(by_group),row.group_mae,atol=1e-14) and np.isclose(errors.mean(),row.row_mae,atol=1e-14))
    checks['reported_group_and_row_errors_match_saved_predictions']=bool(metric_ok)
    fresh=subprocess.run([sys.executable,'-c',
                          "import joblib,sys; b=joblib.load(sys.argv[1]); assert 'pipeline' in b and len(b['columns'])>0",str(OUT/'RF_MODEL.joblib')],
                         capture_output=True,text=True)
    checks['portable_model_loads_in_fresh_python_process']=fresh.returncode==0
    from verify_amendment import verify
    checks.update(verify(ROOT,m,cb.reset_index(),x,receipt,model))
    result={'status':'PASS' if all(checks.values()) else 'FAIL','checks':checks,'passed':sum(checks.values()),'total':len(checks)}
    (OUT/'VALIDATION_CHECKS.json').write_text(json.dumps(result,indent=2)+'\n')
    assert all(checks.values()),result
    report(m,pred,summary,audit,result)
    files={str(p.relative_to(ROOT)):digest(p) for p in ROOT.rglob('*') if p.is_file() and '.venv' not in p.parts and '.mplconfig' not in p.parts and '__pycache__' not in p.parts and '.cache' not in p.parts and p.name!='OUTPUT_HASHES.json'}
    (OUT/'OUTPUT_HASHES.json').write_text(json.dumps(files,indent=2)+'\n')
    print(json.dumps(result,indent=2))


def table(metrics):
    lines=['| Model | Group MAE | Row MAE | RMSE | R² | Class accuracy | Macro-F1 |',
           '|---|---:|---:|---:|---:|---:|---:|']
    for r in metrics:
        lines.append(f"| {r['model']} | {r['group_mae']:.4f} | {r['row_mae']:.4f} | {r['rmse']:.4f} | {r['r2']:.3f} | {r['class_accuracy']:.1%} | {r['class_macro_f1']:.3f} |")
    return '\n'.join(lines)


from report_v2 import report

if __name__ == "__main__":
    main()
